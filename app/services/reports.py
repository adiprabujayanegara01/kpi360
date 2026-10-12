"""Laporan: employee (PDF/CSV), department (PDF/Excel), management (PDF/Excel). Dijalankan sebagai background job."""
import csv
import io
import os
import uuid
from concurrent.futures import ThreadPoolExecutor
from flask import current_app
from ..extensions import db
from ..models import (ReportJob, Employee, EvaluationPeriod, Department, User, Feedback, ADMIN, SUPERVISOR, EMPLOYEE)
from ..analytics import queries as aq
from ..security.rbac import can_access_employee
from ..security.audit import audit
from .common import BusinessError, notify

_pool = ThreadPoolExecutor(max_workers=2)
FORMATS = {"employee": ("pdf", "csv"), "department": ("pdf", "xlsx"), "management": ("pdf", "xlsx")}


def _sections(user, rtype, params):
    p = db.session.get(EvaluationPeriod, int(params["period_id"]))
    if not p:
        raise BusinessError("NOT_FOUND", "Periode tidak ditemukan.", 404)
    if rtype == "employee":
        emp = db.session.get(Employee, int(params["employee_id"]))
        if not emp or not can_access_employee(user, emp.id):
            raise BusinessError("FORBIDDEN", "Tidak berhak atas data karyawan ini.", 403)
        hist = aq.employee_history(emp.id, p.start_date)
        cur = next((h for h in hist if h["period_id"] == p.id), None)
        kp = aq.kpi_breakdown(emp.id, p.id)
        fb = Feedback.query.filter_by(employee_id=emp.id, period_id=p.id).order_by(Feedback.id.desc()).all()
        return (f"Laporan Kinerja Karyawan - {emp.name}", f"Periode {p.name}", [
            ("Data karyawan", ["Field", "Nilai"], [["Nama", emp.name], ["No. karyawan", emp.employee_number],
                                                   ["Departemen", emp.department.name if emp.department else "-"],
                                                   ["Jabatan", emp.position.name if emp.position else "-"],
                                                   ["Skor", f"{cur['score']:.2f}" if cur else "-"], ["Kategori", cur["category"] if cur else "-"]]),
            ("KPI", ["KPI", "Tipe", "Target", "Realisasi", "Achievement %", "Bobot %", "Skor berbobot", "Status"],
             [[k["kpi"], k["type"], k["target"], "-" if k["realization"] is None else k["realization"], round(k["achievement"], 2),
               k["weight"], round(k["weighted"], 2), k["status"]] for k in kp]),
            ("Tren skor", ["Periode", "Skor", "Rata-rata departemen"], [[h["period"], h["score"], h["dept_avg"]] for h in hist]),
            ("Feedback", ["Evaluator", "Komentar", "Rencana pengembangan"], [[f.evaluator.display_name, f.comment, f.development_plan or "-"] for f in fb]),
        ])
    df = aq.score_rows(user, p.id, department_id=params.get("department_id") if rtype == "department" else None)
    r = aq.ranking(df)
    rows = [[int(x.rank), x.employee, x.department or "-", x.position or "-", x.score, x.category, x.incomplete] for x in r.itertuples()]
    head = ["#", "Karyawan", "Departemen", "Jabatan", "Skor", "Kategori", "KPI incomplete"]
    kstats = aq.kpi_stats(user, p.id, params.get("department_id") if rtype == "department" else None)
    ksec = ("Pencapaian per KPI", ["KPI", "Rata-rata achievement %", "Tidak tercapai", "Dari", "Incomplete"],
            [[k["kpi"], k["avg"], k["unmet"], k["n"], k["incomplete"]] for k in kstats])
    if rtype == "department":
        d = db.session.get(Department, int(params["department_id"])) if params.get("department_id") else None
        return (f"Laporan Kinerja Departemen - {d.name if d else 'Semua'}" if user.role == ADMIN else "Laporan Kinerja Tim",
                f"Periode {p.name}", [("Ringkasan", ["Metrik", "Nilai"], [["Jumlah karyawan", len(df)], ["Rata-rata skor", round(float(df.score.mean()), 2) if len(df) else "-"],
                                                                         ["Tertinggi", f"{r.iloc[0].employee} ({r.iloc[0].score})" if len(r) else "-"],
                                                                         ["Terendah", f"{r.iloc[-1].employee} ({r.iloc[-1].score})" if len(r) else "-"]]),
                                       ("Peringkat", head, rows), ksec])
    cards = aq.summary_cards(df)
    narrative = (f"Pada periode {p.name}, rata-rata skor organisasi {cards['avg']} dari {cards['count']} karyawan. "
                 f"Departemen tertinggi: {aq.dept_averages(df)[0]['department']}; terendah: {aq.dept_averages(df)[-1]['department']}. "
                 f"KPI paling sering tidak tercapai: {kstats[0]['kpi'] if kstats else '-'}. (Ringkasan otomatis - perlu ditinjau manusia.)") if cards["count"] else "Belum ada data."
    return ("Laporan Manajemen Kinerja", f"Periode {p.name}", [
        ("Ringkasan", ["Uraian"], [[narrative]]),
        ("Rata-rata per departemen", ["Departemen", "Rata-rata", "Jumlah"], [[x["department"] or "-", x["avg"], x["n"]] for x in aq.dept_averages(df)]),
        ("Distribusi kategori", ["Kategori", "Jumlah"], [[x["category"], x["count"]] for x in aq.distribution(df)]),
        ksec, ("Peringkat", head, rows)])


def render(title, subtitle, sections, fmt):
    buf = io.BytesIO()
    if fmt == "csv":
        s = io.StringIO()
        w = csv.writer(s)
        w.writerow([title, subtitle])
        for h, cols, rows in sections:
            w.writerow([])
            w.writerow([h])
            w.writerow(cols)
            w.writerows(rows)
        buf.write(("\ufeff" + s.getvalue()).encode("utf-8"))
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "Ringkasan"
        ws0["A1"], ws0["A2"] = title, subtitle
        ws0["A1"].font = Font(bold=True, size=14)
        for i, (h, cols, rows) in enumerate(sections):
            ws = wb.create_sheet(h[:28].replace("/", "-"))
            ws.append(cols)
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F6FB2")
            for r in rows:
                ws.append(list(r))
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
            ws0.append([h, f"{len(rows)} baris"])
        wb.save(buf)
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        st = getSampleStyleSheet()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=title, leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
        story = [Paragraph(title, st["Title"]), Paragraph(subtitle, st["Normal"]), Spacer(1, 10)]
        for h, cols, rows in sections:
            story += [Paragraph(h, st["Heading3"])]
            data = [cols] + [[Paragraph(str(c), st["BodyText"]) if len(str(c)) > 30 else str(c) for c in r] for r in rows] if rows else [cols, ["-"] * len(cols)]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6FB2")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                   ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C5CCD6")),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 10)]
        doc.build(story)
    return buf.getvalue()


def _run(app, job_id):
    with app.app_context():
        job = db.session.get(ReportJob, job_id)
        try:
            user = db.session.get(User, job.user_id)
            title, sub, sections = _sections(user, job.report_type, job.params)
            data = render(title, sub, sections, job.fmt)
            os.makedirs(current_app.config["REPORT_FOLDER"], exist_ok=True)
            fname = f"{uuid.uuid4().hex}.{job.fmt}"
            with open(os.path.join(current_app.config["REPORT_FOLDER"], fname), "wb") as f:
                f.write(data)
            job.file_path, job.status = fname, "DONE"
            notify(job.user_id, "REPORT_READY", f"Laporan siap: {title}", sub, f"/reports")
        except Exception as e:  # noqa
            db.session.rollback()
            job = db.session.get(ReportJob, job_id)
            job.status, job.error = "FAILED", str(e)[:390]
            notify(job.user_id, "REPORT_FAILED", "Laporan gagal dibuat", job.error, "/reports")
        db.session.commit()


def create_job(user, rtype, fmt, params):
    if rtype not in FORMATS or fmt not in FORMATS[rtype]:
        raise BusinessError("VALIDATION_ERROR", "Jenis/format laporan tidak valid.", 422)
    if rtype == "management" and user.role != ADMIN:
        raise BusinessError("FORBIDDEN", "Laporan manajemen hanya untuk HR.", 403)
    if rtype == "department" and user.role == EMPLOYEE:
        raise BusinessError("FORBIDDEN", "Tidak berhak.", 403)
    if rtype == "employee":
        eid = params.get("employee_id") or (user.employee.id if user.employee else None)
        params["employee_id"] = eid
        if not eid or not can_access_employee(user, int(eid)):
            raise BusinessError("FORBIDDEN", "Tidak berhak atas data karyawan ini.", 403)
    if not params.get("period_id"):
        raise BusinessError("VALIDATION_ERROR", "Pilih periode.", 422)
    job = ReportJob(user_id=user.id, report_type=rtype, fmt=fmt, params=params)
    db.session.add(job)
    audit("EXPORT_REPORT", "report_jobs", None, new={"type": rtype, "fmt": fmt, **params}, user=user)
    db.session.commit()
    app = current_app._get_current_object()
    if app.config["REPORTS_SYNC"]:
        _run(app, job.id)
        db.session.refresh(job)
    else:
        _pool.submit(_run, app, job.id)
    return job
